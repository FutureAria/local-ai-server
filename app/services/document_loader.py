import importlib.util
import io
import shutil
from html.parser import HTMLParser
from pathlib import Path
from typing import NamedTuple


SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx", ".html", ".htm"}
PDF_OCR_TEXT_THRESHOLD = 20
OCR_INSTALL_HINT = (
    "PDF OCR을 사용하려면 Python dependency와 로컬 tesseract binary가 필요합니다: "
    "pip install -e '.[ocr]' 후 macOS는 'brew install tesseract', "
    "Ubuntu/Debian은 'sudo apt install tesseract-ocr'를 실행하세요. "
    "한국어 OCR은 tesseract language pack(kor)을 별도로 설치해야 합니다."
)
DOCUMENT_TYPE_REQUIREMENTS = {
    ".txt": {
        "file_type": "txt",
        "optional_dependency": None,
        "description": "UTF-8 plain text",
    },
    ".md": {
        "file_type": "md",
        "optional_dependency": None,
        "description": "UTF-8 Markdown text",
    },
    ".pdf": {
        "file_type": "pdf",
        "optional_dependency": "pypdf",
        "description": "Text-based PDF with optional OCR fallback for PyPDF image XObjects.",
    },
    ".docx": {
        "file_type": "docx",
        "optional_dependency": "python-docx",
        "module_name": "docx",
        "description": "Microsoft Word DOCX text and table extraction.",
    },
    ".html": {
        "file_type": "html",
        "optional_dependency": None,
        "description": "UTF-8 HTML text extraction. Script/style content is ignored.",
    },
    ".htm": {
        "file_type": "html",
        "optional_dependency": None,
        "description": "UTF-8 HTML text extraction. Script/style content is ignored.",
    },
}


class DocumentLoaderError(ValueError):
    pass


class OcrPageResult(NamedTuple):
    text: str
    skipped_reason: str | None = None


class _HTMLTextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._parts: list[str] = []
        self._ignored_depth = 0

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag.lower() in {"head", "script", "style", "noscript"}:
            self._ignored_depth += 1
        if tag.lower() in {"p", "br", "div", "section", "article", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6"}:
            self._parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() in {"head", "script", "style", "noscript"} and self._ignored_depth:
            self._ignored_depth -= 1
        if tag.lower() in {"p", "div", "section", "article", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6"}:
            self._parts.append("\n")

    def handle_data(self, data: str) -> None:
        if self._ignored_depth:
            return
        text = data.strip()
        if text:
            self._parts.append(text)

    def text(self) -> str:
        lines = [" ".join(line.split()) for line in "".join(self._parts).splitlines()]
        return "\n".join(line for line in lines if line).strip()


def file_type_for_path(path: str | Path) -> str:
    suffix = Path(path).suffix.lower()
    requirement = DOCUMENT_TYPE_REQUIREMENTS.get(suffix)
    if requirement is None:
        return suffix.lstrip(".")
    return requirement["file_type"]


class DocumentLoader:
    def _module_available(self, module_name: str) -> bool:
        try:
            return importlib.util.find_spec(module_name) is not None
        except (ImportError, ValueError):
            return False

    def pdf_ocr_status(self) -> dict:
        missing = []
        if not self._module_available("pytesseract"):
            missing.append("pytesseract")
        if not self._module_available("PIL"):
            missing.append("Pillow")
        if shutil.which("tesseract") is None:
            missing.append("tesseract binary")
        return {
            "available": not missing,
            "missing": missing,
            "install_hint": None if not missing else OCR_INSTALL_HINT,
        }

    def supported_types(self) -> list[dict]:
        types = []
        for extension in sorted(SUPPORTED_EXTENSIONS):
            requirement = DOCUMENT_TYPE_REQUIREMENTS[extension]
            module_name = requirement.get("module_name") or requirement["optional_dependency"]
            optional_dependency = requirement["optional_dependency"]
            available = optional_dependency is None or self._module_available(module_name)
            types.append(
                {
                    "extension": extension,
                    "file_type": requirement["file_type"],
                    "available": available,
                    "optional_dependency": optional_dependency,
                    "install_hint": None if available else "pip install -e '.[documents]'",
                    "description": requirement["description"],
                }
            )
        return types

    def load_text(self, path: str | Path) -> str:
        file_path = Path(path)
        suffix = file_path.suffix.lower()
        if suffix not in SUPPORTED_EXTENSIONS:
            raise DocumentLoaderError(
                f"지원하지 않는 파일 형식입니다: {suffix}. 지원 형식: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
            )
        if not file_path.exists() or not file_path.is_file():
            raise DocumentLoaderError(f"파일을 찾을 수 없습니다: {file_path}")
        if suffix == ".pdf":
            return self._load_pdf(file_path)
        if suffix == ".docx":
            return self._load_docx(file_path)
        if suffix in {".html", ".htm"}:
            return self._load_html(file_path)
        try:
            return file_path.read_text(encoding="utf-8")
        except UnicodeDecodeError as exc:
            raise DocumentLoaderError(f"UTF-8 텍스트 파일만 지원합니다: {file_path}") from exc

    def _load_pdf(self, file_path: Path) -> str:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise DocumentLoaderError(
                "PDF 문서를 읽으려면 optional dependency가 필요합니다: "
                "pip install -e '.[documents]'"
            ) from exc

        try:
            reader = PdfReader(str(file_path))
            pages = list(reader.pages)
        except Exception as exc:
            raise DocumentLoaderError(f"PDF 텍스트 추출에 실패했습니다: {file_path}") from exc

        parts: list[str] = []
        skipped_ocr_reasons: list[str] = []
        for page_index, page in enumerate(pages, start=1):
            page_text = (page.extract_text() or "").strip()
            if len(page_text) >= PDF_OCR_TEXT_THRESHOLD:
                parts.append(page_text)
                continue

            ocr_result = self._ocr_pdf_page_images(page, page_index)
            if ocr_result.text:
                parts.append(ocr_result.text)
            elif page_text:
                parts.append(page_text)

            if ocr_result.skipped_reason:
                skipped_ocr_reasons.append(ocr_result.skipped_reason)

        text = "\n\n".join(part for part in parts if part.strip()).strip()
        if not text:
            reason = "; ".join(skipped_ocr_reasons) if skipped_ocr_reasons else "OCR 대상 이미지가 없습니다."
            raise DocumentLoaderError(f"PDF에서 추출 가능한 텍스트가 없습니다: {file_path}. OCR fallback skipped: {reason}")
        return text

    def _ocr_pdf_page_images(self, page, page_index: int) -> OcrPageResult:
        images = list(getattr(page, "images", []) or [])
        if not images:
            return OcrPageResult("", f"page {page_index}: PyPDF로 추출 가능한 image XObject가 없습니다.")

        status = self.pdf_ocr_status()
        if not status["available"]:
            return OcrPageResult("", f"page {page_index}: {status['install_hint']}")

        try:
            import pytesseract
            from PIL import Image
        except ImportError:
            return OcrPageResult("", f"page {page_index}: {OCR_INSTALL_HINT}")

        texts: list[str] = []
        skipped: list[str] = []
        for image_index, image in enumerate(images, start=1):
            try:
                data = getattr(image, "data", image)
                with Image.open(io.BytesIO(data)) as pil_image:
                    text = pytesseract.image_to_string(pil_image).strip()
            except Exception as exc:
                skipped.append(f"page {page_index} image {image_index}: OCR 실패({exc.__class__.__name__})")
                continue
            if text:
                texts.append(text)

        if texts:
            return OcrPageResult("\n".join(texts).strip(), None)
        reason = "; ".join(skipped) if skipped else f"page {page_index}: OCR 결과 텍스트가 비어 있습니다."
        return OcrPageResult("", reason)

    def _load_docx(self, file_path: Path) -> str:
        try:
            from docx import Document as DocxDocument
        except ImportError as exc:
            raise DocumentLoaderError(
                "DOCX 문서를 읽으려면 optional dependency가 필요합니다: "
                "pip install -e '.[documents]'"
            ) from exc

        try:
            document = DocxDocument(str(file_path))
            parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
        except Exception as exc:
            raise DocumentLoaderError(f"DOCX 텍스트 추출에 실패했습니다: {file_path}") from exc

        text = "\n".join(parts).strip()
        if not text:
            raise DocumentLoaderError(f"DOCX에서 추출 가능한 텍스트가 없습니다: {file_path}")
        return text

    def _load_html(self, file_path: Path) -> str:
        try:
            html = file_path.read_text(encoding="utf-8")
        except UnicodeDecodeError as exc:
            raise DocumentLoaderError(f"UTF-8 HTML 파일만 지원합니다: {file_path}") from exc

        parser = _HTMLTextExtractor()
        try:
            parser.feed(html)
            parser.close()
        except Exception as exc:
            raise DocumentLoaderError(f"HTML 텍스트 추출에 실패했습니다: {file_path}") from exc

        text = parser.text()
        if not text:
            raise DocumentLoaderError(f"HTML에서 추출 가능한 텍스트가 없습니다: {file_path}")
        return text
